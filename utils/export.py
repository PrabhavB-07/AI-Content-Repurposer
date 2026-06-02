import io
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ── Section parser ─────────────────────────────────────────────────────────────
SECTIONS = [
    ("📸 Instagram Post",  "INSTAGRAM:"),
    ("💼 LinkedIn Post",   "LINKEDIN:"),
    ("🐦 Twitter Thread",  "TWITTER:"),
    ("📝 Blog Post",       "BLOG:"),
]

def parse_sections(raw: str) -> list[dict]:
    """Return list of {title, body} dicts parsed from raw AI output."""
    import re
    results = []
    keys = [s[1] for s in SECTIONS]

    for i, (title, key) in enumerate(SECTIONS):
        next_key = keys[i + 1] if i + 1 < len(keys) else None
        if next_key:
            pattern = rf"{re.escape(key)}(.*?){re.escape(next_key)}"
        else:
            pattern = rf"{re.escape(key)}(.*)"
        match = re.search(pattern, raw, re.DOTALL)
        body = match.group(1).strip() if match else ""
        results.append({"title": title, "body": body})

    return results


# ── PDF Export ─────────────────────────────────────────────────────────────────
def generate_pdf(raw_content: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2*cm, rightMargin=2*cm,
        topMargin=2.5*cm, bottomMargin=2*cm,
    )

    styles = getSampleStyleSheet()

    style_title = ParagraphStyle(
        "AppTitle",
        fontName="Helvetica-Bold",
        fontSize=22,
        textColor=colors.HexColor("#3b6ef8"),
        spaceAfter=4,
        alignment=TA_CENTER,
    )
    style_subtitle = ParagraphStyle(
        "AppSubtitle",
        fontName="Helvetica",
        fontSize=10,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=16,
        alignment=TA_CENTER,
    )
    style_section = ParagraphStyle(
        "SectionHead",
        fontName="Helvetica-Bold",
        fontSize=13,
        textColor=colors.HexColor("#1e293b"),
        spaceBefore=18,
        spaceAfter=6,
    )
    style_body = ParagraphStyle(
        "Body",
        fontName="Helvetica",
        fontSize=10,
        textColor=colors.HexColor("#334155"),
        leading=16,
        spaceAfter=4,
    )

    story = []

    # Header
    story.append(Paragraph("AI Content Repurposer", style_title))
    story.append(Paragraph("Generated with Groq AI", style_subtitle))
    story.append(HRFlowable(width="100%", thickness=1,
                             color=colors.HexColor("#e2e8f0"), spaceAfter=12))

    # Sections
    sections = parse_sections(raw_content)
    for sec in sections:
        if not sec["body"]:
            continue
        story.append(Paragraph(sec["title"], style_section))
        story.append(HRFlowable(width="100%", thickness=0.5,
                                 color=colors.HexColor("#e2e8f0"), spaceAfter=6))
        # Split into paragraphs on blank lines
        for para in sec["body"].split("\n\n"):
            para = para.strip()
            if para:
                # Escape XML special chars for ReportLab
                para = para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(para.replace("\n", "<br/>"), style_body))
                story.append(Spacer(1, 4))
        story.append(Spacer(1, 8))

    doc.build(story)
    buffer.seek(0)
    return buffer.read()


# ── DOCX Export ────────────────────────────────────────────────────────────────
def generate_docx(raw_content: str) -> bytes:
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.left_margin   = Inches(1)
        section.right_margin  = Inches(1)
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)

    # App title
    title = doc.add_heading("AI Content Repurposer", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x3b, 0x6e, 0xf8)
        run.font.size = Pt(22)

    sub = doc.add_paragraph("Generated with Groq AI")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
    sub.runs[0].font.size = Pt(10)

    doc.add_paragraph()  # spacer

    # Sections
    sections = parse_sections(raw_content)
    for sec in sections:
        if not sec["body"]:
            continue

        # Section heading
        h = doc.add_heading(sec["title"], level=2)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
            run.font.size = Pt(13)

        # Body paragraphs
        for para in sec["body"].split("\n\n"):
            para = para.strip()
            if para:
                p = doc.add_paragraph(para)
                p.paragraph_format.space_after = Pt(6)
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        doc.add_paragraph()  # spacer between sections

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()